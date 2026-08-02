from __future__ import annotations

import json
import math
from collections import Counter, defaultdict
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "merge_plan_data.json"
OUTPUT_DIR = ROOT.parent.parent / "outputs" / "merge_handoff_20260801"
QA_DIR = ROOT.parent.parent / "qa" / "merge_handoff_20260801"
OUTPUT_PATH = OUTPUT_DIR / "birdpart2_项目合并实施书_2026-08-01.docx"
ARCH_PATH = ROOT / "birdpart2_merge_architecture.png"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GREEN = "E2F0D9"
PALE_AMBER = "FFF2CC"
PALE_RED = "FCE4D6"
PALE_GRAY = "F2F4F7"
MID_GRAY = "D9E2F3"
TEXT = "202124"
MUTED = "5F6B76"
WHITE = "FFFFFF"


def load_data() -> dict:
    return json.loads(DATA_PATH.read_text(encoding="utf-8"))


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C3D0", size=6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_table_indent(table, twips=120) -> None:
    tbl_pr = table._tbl.tblPr
    ind = tbl_pr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tbl_pr.append(ind)
    ind.set(qn("w:w"), str(twips))
    ind.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_fixed_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeatable_font(run, size=11, bold=False, color=TEXT, name="Microsoft YaHei") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), name)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_char2])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Title", 24, NAVY, 0, 10),
        ("Subtitle", 12, MUTED, 0, 12),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = doc.styles[list_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    if "Caption Custom" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("Caption Custom", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["Caption Custom"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(9)
    style.font.color.rgb = RGBColor.from_string(MUTED)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_before = Pt(3)
    style.paragraph_format.space_after = Pt(9)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("BIRDPART2  ·  项目合并实施书  ·  受控基线")
    set_repeatable_font(run, size=8.5, bold=True, color=DARK_BLUE)
    p.paragraph_format.space_after = Pt(0)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("birdbox-v1@1.0.0  ·  第 ")
    set_repeatable_font(run, size=8.5, color=MUTED)
    add_field(p, "PAGE")
    run = p.add_run(" 页")
    set_repeatable_font(run, size=8.5, color=MUTED)


def add_paragraph(doc, text: str = "", style=None, bold_prefix: str | None = None, color=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_repeatable_font(r1, bold=True, color=color or TEXT)
        r2 = p.add_run(text[len(bold_prefix):])
        set_repeatable_font(r2, color=color or TEXT)
    else:
        r = p.add_run(text)
        set_repeatable_font(r, color=color or TEXT)
    return p


def add_bullet(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_repeatable_font(r)


def add_number(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_repeatable_font(r)


def add_table(doc, headers, rows, widths, font_size=8.8, header_color=NAVY, alignments=None):
    if not math.isclose(sum(widths), 6.5, abs_tol=0.05):
        raise ValueError(f"Table widths must total 6.5 inches: {widths}")
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_indent(table, 120)
    set_table_borders(table)
    header = table.rows[0]
    repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_fixed_width(cell, widths[i])
        set_cell_shading(cell, header_color)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(value))
        set_repeatable_font(r, size=font_size, bold=True, color=WHITE)
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        if row_idx % 2 == 1:
            for cell in row.cells:
                set_cell_shading(cell, "F8FAFC")
        for i, value in enumerate(values):
            cell = row.cells[i]
            set_fixed_width(cell, widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            if alignments and i < len(alignments):
                p.alignment = alignments[i]
            r = p.add_run("" if value is None else str(value))
            set_repeatable_font(r, size=font_size, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_status_cell_color(table, col_idx: int) -> None:
    colors = {
        "已完成": PALE_GREEN,
        "进行中": PALE_AMBER,
        "阻塞": PALE_RED,
        "待开始": PALE_GRAY,
        "已接入": PALE_GREEN,
        "客户端已接入": PALE_AMBER,
        "客户端与 mock 已接入": PALE_AMBER,
        "部分一致": PALE_RED,
    }
    for row in table.rows[1:]:
        text = row.cells[col_idx].text.strip()
        if text in colors:
            set_cell_shading(row.cells[col_idx], colors[text])


def add_callout(doc, title: str, body: str, fill=PALE_BLUE, accent=BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_indent(table, 120)
    set_table_borders(table, color=accent, size=8)
    cell = table.cell(0, 0)
    set_fixed_width(cell, 6.5)
    set_cell_margins(cell, top=120, start=160, bottom=120, end=160)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_repeatable_font(r, size=10.5, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.08
    r2 = p2.add_run(body)
    set_repeatable_font(r2, size=10, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def draw_architecture() -> None:
    width, height = 1900, 970
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    bold_path = Path(r"C:\Windows\Fonts\msyhbd.ttc")
    if not font_path.exists():
        font_path = Path(r"C:\Windows\Fonts\simsun.ttc")
    if not bold_path.exists():
        bold_path = font_path
    title_font = ImageFont.truetype(str(bold_path), 54)
    box_title_font = ImageFont.truetype(str(bold_path), 34)
    body_font = ImageFont.truetype(str(font_path), 28)
    footer_font = ImageFont.truetype(str(bold_path), 32)
    note_font = ImageFont.truetype(str(font_path), 27)

    def rgb(value: str):
        return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))

    def box(x, y, w, h, title, lines, fill, edge=BLUE):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=18, fill=rgb(fill), outline=rgb(edge), width=4)
        draw.text((x + 24, y + 20), title, font=box_title_font, fill=rgb(DARK_BLUE))
        draw.multiline_text((x + 24, y + 76), "\n".join(lines), font=body_font, fill=rgb(TEXT), spacing=14)

    def arrow(x1, y1, x2, y2, color=BLUE):
        draw.line((x1, y1, x2, y2), fill=rgb(color), width=5)
        head = 18
        draw.polygon([(x2, y2), (x2 - head, y2 - head // 2), (x2 - head, y2 + head // 2)], fill=rgb(color))

    draw.text((34, 28), "birdpart2 目标运行时与业务对接链路", font=title_font, fill=rgb(NAVY))
    top = [
        (34, "生产入口", ["lib/main.dart", "IntegratedBirdBootstrap"]),
        (410, "统一 Shell", ["相册 · 处理进度 · 我的", "三组嵌套 Navigator"]),
        (786, "应用协调", ["RepositoryTaskController", "Cubits · SettingsController"]),
        (1162, "Repository", ["Storage / Batch / Job", "Photo / Review / Copy"]),
        (1538, "盒子服务", ["REST /api/v1", "WebSocket /events"]),
    ]
    box_w, box_h, box_y = 326, 230, 160
    for i, (x, title, lines) in enumerate(top):
        box(x, box_y, box_w, box_h, title, lines, PALE_BLUE if i < 4 else PALE_GREEN)
        if i < len(top) - 1:
            arrow(x + box_w + 6, box_y + box_h // 2, top[i + 1][0] - 10, box_y + box_h // 2)

    box(34, 500, 520, 220, "共享核心", ["ApiClient · EventClient · SessionCoordinator", "Hive · BirdSyncService · AppDataChangeBus"], PALE_GRAY, edge=DARK_BLUE)
    box(690, 500, 520, 220, "贯通主键", ["device_id → project_id", "file_id / job_id / version / idempotency key"], PALE_AMBER, edge="C58A00")
    box(1346, 500, 520, 220, "发布边界", ["mock：开发与契约证据", "真实盒子 + 签名：正式发布证据"], PALE_RED, edge="C65911")
    arrow(564, 610, 680, 610, MUTED)
    arrow(1220, 610, 1336, 610, MUTED)

    draw.text((34, 790), "业务流：连接/配对 → SD 扫描/建项目 → 导入/分析 → 相册/审阅 → 复制/任务/报告/日志", font=footer_font, fill=rgb(NAVY))
    draw.text((34, 858), "关键原则：一个用户动作只经过一条生产路径；mock 通过不等于真实盒子验收通过。", font=note_font, fill=rgb(MUTED))
    image.save(ARCH_PATH, format="PNG", optimize=True)


def batch_metrics(data):
    grouped = defaultdict(list)
    for task in data["tasks"]:
        grouped[task["batch"]].append(task)
    metrics = {}
    for batch, tasks in grouped.items():
        weight = sum(float(t["weight"]) for t in tasks)
        earned = sum(float(t["weight"]) * float(t["progress"]) for t in tasks)
        metrics[batch] = {
            "weight": weight,
            "progress": earned / weight if weight else 0,
            "counts": Counter(t["status"] for t in tasks),
        }
    return metrics


def build_docx(data: dict) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)
    draw_architecture()

    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "birdpart2 项目合并实施书"
    props.subject = "Aves 与 Part2 合并实施与进度基线"
    props.author = "项目集成组"
    props.comments = "基于 2026-08-01 当前工作区、双方交接资料和 BirdBox v1 契约生成。"

    meta = data["meta"]
    tasks = data["tasks"]
    total_weight = sum(float(t["weight"]) for t in tasks)
    earned = sum(float(t["weight"]) * float(t["progress"]) for t in tasks)
    overall = earned / total_weight
    counts = Counter(t["status"] for t in tasks)
    metrics = batch_metrics(data)

    # Cover using the memo_masthead template with first-page header suppressed.
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("BIRDPART2  ·  MERGE DELIVERY BASELINE")
    set_repeatable_font(r, size=9, bold=True, color=BLUE)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("项目合并实施书")
    set_repeatable_font(r, size=24, bold=True, color=NAVY)
    p = doc.add_paragraph(style="Subtitle")
    r = p.add_run("Aves × Part2 运行时整合、功能对接、剩余实施与发布门禁")
    set_repeatable_font(r, size=12, color=MUTED)

    add_table(
        doc,
        ["项目", "当前基线"],
        [
            ["工作区", meta["project_path"]],
            ["分支 / HEAD", f"{meta['branch']} / {meta['head'][:12]}"],
            ["契约", meta["contract"]],
            ["基线日期", meta["baseline_date"]],
            ["建议完成日期", f"{meta['proposed_finish']}（需四方确认）"],
            ["文档版本", meta["document_version"]],
        ],
        [1.55, 4.95],
        font_size=9.3,
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    add_callout(
        doc,
        "执行判定",
        f"Git 层和客户端运行时合并已基本完成，统一入口、三 Tab、Repository 任务链路、相册审阅、复制、日志和设置均已进入同一依赖图。按 43 项加权任务计算，当前总体进度为 {overall:.1%}。但真实盒子基线、真机端到端和正式签名未闭环，因此当前结论是“可继续联调，不可正式发布”。",
        fill=PALE_AMBER,
        accent="C58A00",
    )
    add_paragraph(doc, "保密与使用说明：本文不包含真实 Token、密钥或数据库账号；建议排期不替代项目负责人确认。", color=MUTED)
    doc.add_page_break()

    # Executive summary.
    doc.add_heading("执行摘要", level=1)
    add_paragraph(
        doc,
        "birdpart2 不是简单把两个目录合在一起，而是把 Aves 提供的连接、相册、审阅、离线同步、复制和任务 Repository，与 Part2 提供的任务、设备和设置体验，收敛到同一个生产入口、同一份依赖容器和同一套 BirdBox v1 契约。当前代码已经完成主要客户端接线，剩余工作集中在协议尾项清理、当前树复测、真实盒子联调和发布证据。",
    )

    add_table(
        doc,
        ["指标", "当前值", "解释"],
        [
            ["加权总体进度", f"{overall:.2%}", "包括客户端、真盒、签名和发布门禁"],
            ["任务数", str(len(tasks)), f"已完成 {counts['已完成']}；进行中 {counts['进行中']}；阻塞 {counts['阻塞']}；待开始 {counts['待开始']}"],
            ["当前工作区", "155 项变化", "107 修改、47 未跟踪、1 删除；以 live tree 为准"],
            ["当前契约", meta["contract"], "非严格校验 2026-08-01 通过；strict-baseline 缺 13 项"],
            ["历史 B7 证据", "Preflight passed", "2026-07-29；脏工作树；Debug APK 202,368,038 bytes"],
            ["发布结论", "NO-GO", "真实盒子、正式签名、当前树完整复跑未满足"],
        ],
        [1.45, 1.55, 3.5],
        font_size=9,
    )

    doc.add_heading("需要优先关闭的 8 项", level=2)
    for task_id in ("T04", "T14", "T15", "T35", "T36", "T40", "T41", "T42"):
        task = next(t for t in tasks if t["id"] == task_id)
        add_bullet(doc, f"{task_id}｜{task['task']}（{task['status']}，owner：{task['owner']}，目标：{task['planned_end']}）")

    doc.add_heading("1. 项目背景、范围与完成定义", level=1)
    doc.add_heading("1.1 两个源项目的责任", level=2)
    add_table(
        doc,
        ["来源", "核心能力", "在合并后的定位"],
        [
            ["Aves", "连接、相册、照片/场景/连拍、详情/对比/审阅、离线同步、复制与任务 Repository、mock", "数据与业务能力底座；负责真实照片链路"],
            ["Part2", "任务首页/详情/结果、SD 卡流程、设备/设置/诊断体验", "任务与设备交互层；改由真实 Repository 驱动"],
            ["BirdBox", "设备、文件、数据库、任务执行、媒体、日志和版本仲裁", "权威后端；本仓库没有真实盒子源码"],
        ],
        [1.0, 2.6, 2.9],
        font_size=9,
    )
    doc.add_heading("1.2 合并目标", level=2)
    for text in (
        "默认 lib/main.dart 只启动一个生产应用树；demo 只能显式运行。",
        "全局只创建一套 ApiClient、EventClient、SessionCoordinator、Hive 缓存和 Repository。",
        "device_id → project_id → file_id/job_id 贯穿任务、相册、审阅、复制与报告。",
        "HTTP、WebSocket、媒体和日志遵循同一配对与鉴权策略。",
        "离线写复用幂等键；version 冲突返回 409，App 不静默覆盖。",
        "mock 用于开发与契约验证；正式完成必须有真实盒子、签名和端到端证据。",
    ):
        add_bullet(doc, text)
    doc.add_heading("1.3 本轮不在 App 仓库内完成的事项", level=2)
    for text in (
        "真实盒子文件扫描、AI 推理、复制执行、任务调度和服务部署。",
        "真实盒子数据库 DDL、迁移、回滚、备份和多客户端锁实现。",
        "正式 keystore 的生成、保管和密码分发；本文只定义注入与验签门。",
        "云账号、云存储、支付、NAS 复杂同步或 BLE 完整配网。",
    ):
        add_bullet(doc, text)

    doc.add_heading("2. 当前项目结构与功能对接", level=1)
    add_paragraph(doc, "当前生产入口为 lib/main.dart → runIntegratedBirdApp()。IntegratedBirdBootstrap 创建唯一 BirdCompanionDependencies，并由 BirdCompanionScope 向三 Tab 和路由页面注入。页面不直接发 HTTP，而是通过 Controller/Cubit → Repository → Api → ApiClient 调用盒子。")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(ARCH_PATH), width=Inches(6.45))
    cap = doc.add_paragraph("图 1  birdpart2 统一运行时与端到端业务链路", style="Caption Custom")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("2.1 功能对接矩阵", level=2)
    interface_rows = []
    for item in data["interfaces"]:
        interface_rows.append([
            f"{item['id']}\n{item['flow']}",
            f"{item['source_module']}\n→ {item['target_module']}",
            f"{item['mechanism']}\n协议：{item['protocol']}",
            f"{item['status']}\n{item['gap']}",
        ])
    table = add_table(doc, ["链路", "模块对接", "怎么对接", "当前结论"], interface_rows, [1.05, 2.0, 2.15, 1.3], font_size=7.9)
    add_status_cell_color(table, 3)

    doc.add_heading("2.2 模块间同步原则", level=2)
    for text in (
        "设备会话切换由 DeviceSessionCubit 和 SessionCoordinator 管理，REST、WS 和缓存命名空间同步切换。",
        "任务状态由 BirdJobStatus/JobReport 作为权威模型；WS 提供完整快照，断线时运行中任务每 4 秒轮询。",
        "任务分析完成后以 source_project_id 构造 GalleryArgs，切到相册 Tab 后再 push 真实 Gallery 路由。",
        "审阅保存通过 UserDecisionPatch，离线队列只覆盖单张审阅和批量照片操作；任务控制与复制创建不离线排队。",
        "AppDataChangeBus 用精确资源事件协调相册、任务、设备和设置，避免整 App 重建。",
    ):
        add_bullet(doc, text)

    doc.add_heading("3. 当前进度与状态判定", level=1)
    batch_rows = []
    batch_by_id = {b["batch"]: b for b in data["batches"]}
    for batch in data["batches"]:
        m = metrics[batch["batch"]]
        counts_text = " / ".join(f"{key}{m['counts'].get(key, 0)}" for key in ("已完成", "进行中", "阻塞", "待开始"))
        if m["counts"].get("阻塞", 0):
            verdict = "客户端已推进；外部验收阻塞"
        elif m["progress"] >= 0.999:
            verdict = "已完成"
        else:
            verdict = "进行中"
        batch_rows.append([batch["batch"], batch["name"], f"{m['progress']:.1%}", counts_text, verdict])
    table = add_table(doc, ["批次", "主题", "加权进度", "任务状态", "判定"], batch_rows, [0.55, 2.15, 0.9, 1.55, 1.35], font_size=8.4)
    add_status_cell_color(table, 4)
    add_callout(
        doc,
        "进度口径",
        "权重总计 100。代码存在不等于批次完成：涉及真实盒子的 B3-B6 必须有真盒 E2E；B7 必须同时满足 strict baseline、11 项真盒证据、正式签名验签和干净工作区。",
        fill=PALE_BLUE,
        accent=BLUE,
    )

    doc.add_heading("4. B0-B7 分批实施与退出门", level=1)
    exit_gates = {
        "B0": "默认契约校验与 strict-baseline 均通过，四方联系人和真盒基线完整。",
        "B1": "默认入口只有一套生产依赖和三 Tab；无占位相册、无生产 demo ID。",
        "B2": "DTO/错误/迁移/审阅 Patch/多设备隔离测试通过，wire 不再发送 batch_id。",
        "B3": "真盒配对、会话恢复/吊销、REST、WS、媒体与日志鉴权全部通过。",
        "B4": "真盒扫描→建项目→导入→分析→杀 App→恢复同一 job_id 完整通过。",
        "B5": "真盒相册/详情/连拍/对比/审阅、离线重放、409 与双设备隔离通过。",
        "B6": "真盒复制、失败项、报告、日志与设置重启闭环通过。",
        "B7": "当前树 Preflight、strict baseline、11 项真盒 E2E、签名 release 与验签全部通过。",
    }
    rollback = {
        "B0": "整体 revert 契约/文档提交；不得用假值绕过 strict baseline。",
        "B1": "单独 revert 入口/Shell 提交；demo 入口仅作开发验证。",
        "B2": "使用新键前缀与 quarantine；回退时旧 pending 数据仍可读。",
        "B3": "清除内存与 Keystore 会话；回退后不得以关闭鉴权冒充完成。",
        "B4": "保留 demo 数据源作显式开发开关；生产接口失败必须显示错误。",
        "B5": "保留旧缓存只读兼容；不得删除其他设备 pending_operations。",
        "B6": "按复制/任务、日志、设置三个提交组独立回退。",
        "B7": "保留上一稳定 APK、证书指纹和 App/盒子/契约版本；NO-GO 时不分发。",
    }
    tasks_by_batch = defaultdict(list)
    for task in tasks:
        tasks_by_batch[task["batch"]].append(task)

    for batch in data["batches"]:
        batch_id = batch["batch"]
        doc.add_heading(f"4.{int(batch_id[1:]) + 1} {batch_id}｜{batch['name']}", level=2)
        add_paragraph(doc, f"目标：{batch['goal']}。主要责任：{batch['owner']}。", bold_prefix="目标：")
        rows = []
        for task in tasks_by_batch[batch_id]:
            evidence = task["evidence"]
            if len(evidence) > 72:
                evidence = evidence[:69] + "…"
            gap = task["gap"]
            if len(gap) > 62:
                gap = gap[:59] + "…"
            rows.append([
                task["id"],
                task["task"],
                task["status"],
                f"{float(task['progress']):.0%}",
                f"{evidence}\n差距：{gap}",
            ])
        table = add_table(doc, ["ID", "工作项", "状态", "进度", "证据 / 差距"], rows, [0.45, 2.65, 0.75, 0.55, 2.1], font_size=7.8)
        add_status_cell_color(table, 2)
        add_callout(doc, "退出门", exit_gates[batch_id], fill=PALE_GREEN if metrics[batch_id]["progress"] >= 0.99 else PALE_AMBER, accent="548235" if metrics[batch_id]["progress"] >= 0.99 else "C58A00")
        add_paragraph(doc, f"回退策略：{rollback[batch_id]}", bold_prefix="回退策略：", color=MUTED)

    doc.add_heading("5. 建议排期与里程碑", level=1)
    add_paragraph(doc, "以下日期是从 2026-08-01 起算的建议推进节奏，实际日期需在 T04 补齐联系人和联调窗口后冻结。")
    milestone_rows = []
    for m in data["milestones"]:
        milestone_rows.append([m["id"], m["name"], m["planned"], f"{float(m['progress']):.0%}", m["status"], m["criteria"]])
    table = add_table(doc, ["里程碑", "名称", "计划日", "进度", "状态", "完成条件"], milestone_rows, [0.55, 1.55, 0.85, 0.55, 0.7, 2.3], font_size=7.9)
    add_status_cell_color(table, 4)

    doc.add_heading("5.1 两周实施顺序", level=2)
    for text in (
        "2026-08-01—08-04：补真盒基线与联系人；修 T14/T35；释放 Flutter 锁并完成当前树 Preflight。",
        "2026-08-05—08-07：冻结真实鉴权实现，完成 B3 真盒配对/会话/WS/媒体/日志验收。",
        "2026-08-07—08-11：完成 B4 任务恢复与 B5 相册审阅、离线、409 和双设备验收。",
        "2026-08-11—08-13：完成 B6 复制、失败项、报告、日志和设置重启验收。",
        "2026-08-14—08-15：收齐 11 项 E2E 证据，注入签名，执行 Release gate 和 GO/NO-GO。",
    ):
        add_number(doc, text)

    doc.add_heading("6. 风险、问题与依赖", level=1)
    risk_rows = []
    for risk in data["risks"]:
        score = int(risk["probability"]) * int(risk["impact"])
        level = "极高" if score >= 20 else "高" if score >= 12 else "中" if score >= 6 else "低"
        risk_rows.append([
            risk["id"],
            f"{risk['risk']}\n触发：{risk['trigger']}",
            f"{risk['probability']}×{risk['impact']}={score}\n{level}",
            f"{risk['owner']}\n截止：{risk['due']}",
            risk["mitigation"],
        ])
    add_table(doc, ["ID", "风险 / 触发", "评分", "Owner / 截止", "缓解措施"], risk_rows, [0.45, 2.05, 0.7, 1.25, 2.05], font_size=7.7)

    doc.add_heading("7. 质量、联调与发布验收", level=1)
    doc.add_heading("7.1 当前可确认的证据", level=2)
    evidence_rows = [
        ["非严格契约校验", "2026-08-01", "通过", "OpenAPI、refs、schema、fixture 有效"],
        ["strict-baseline", "2026-08-01", "失败（预期）", "缺 13 项真实盒子/联系人字段"],
        ["生产完整性校验", "2026-08-01", "通过", "入口、依赖、fallback 合法"],
        ["B7 Preflight", "2026-07-29", "通过（历史）", "format/analyze/test/debug APK passed；脏工作树"],
        ["当前树 Flutter 全量复跑", "2026-08-01", "未完成", "启动等待无输出后终止，不能声称新鲜通过"],
        ["真实盒子 E2E", "未提供", "阻塞", "不能用 mock 替代"],
        ["正式 release 验签", "未提供", "阻塞", "缺正式签名与批准证书指纹"],
    ]
    table = add_table(doc, ["证据", "日期", "结果", "说明"], evidence_rows, [1.55, 0.9, 1.1, 2.95], font_size=8.5)
    add_status_cell_color(table, 2)

    doc.add_heading("7.2 正式发布必须通过的端到端用例", level=2)
    for text in (
        "mDNS、二维码、手填三种连接及降级路径。",
        "配对、重启恢复、Token 吊销和重新配对。",
        "设备、电量、温度、存储、SD 卡、软件和模型版本。",
        "扫描、创建 project_id、导入/分析 job_id、杀 App 后恢复。",
        "相册、详情、场景、连拍、对比和审阅。",
        "离线写重放与一次 409 人工冲突处理。",
        "keep/all/dual 复制，目标离线、空间不足、正常完成、失败项和报告。",
        "日志导出下载、签名 URL 过期刷新。",
        "正式签名、安装、冷启动、返回键、切设备、缓存升级/降级和回滚。",
    ):
        add_bullet(doc, text)

    doc.add_heading("7.3 GO / NO-GO 判定", level=2)
    add_table(
        doc,
        ["判定", "条件"],
        [
            ["GO", "T01-T43 全部满足；无 P0/P1 未解决；真盒 E2E 通过；正式签名与证书指纹验签通过。"],
            ["NO-GO", "任一 demo/占位路径进入生产；鉴权通道未闭环；离线/409 未验；真盒或签名缺失；工作区与证据不一致。"],
            ["有条件内测", "仅可明确标记测试固件/测试包、关闭生产分发并登记限制；不得命名为正式版。"],
        ],
        [1.1, 5.4],
        font_size=9.2,
    )

    doc.add_heading("8. 组织、变更与回滚管理", level=1)
    add_table(
        doc,
        ["角色", "主要责任", "当前要求"],
        [
            ["协议 owner", "批准 OpenAPI、Schema、错误码和兼容窗口", "补联系人；确认 strict baseline"],
            ["公共 Flutter 核心 owner", "入口、依赖、DTO、缓存、会话和公共路由", "关闭 T14/T35/T36"],
            ["盒子 API owner", "服务实现、固件、数据库、部署与真盒环境", "补 T04 并支持 B3-B6"],
            ["QA / 发布 owner", "同契约测试、真盒证据、签名、构建与发布判定", "完成 T40-T43"],
        ],
        [1.35, 3.15, 2.0],
        font_size=8.7,
    )
    doc.add_heading("8.1 提交与证据规则", level=2)
    for text in (
        "按 B0-B7 拆分契约/模型、实现、测试、文档提交；公共文件指定单一合并人。",
        "不以 git reset --hard 覆盖共享工作；回退使用可审计 revert 或回退分支。",
        "每批记录 App SHA、盒子 SHA、契约/API/schema、命令、结果、APK SHA-256、设备型号和回滚点。",
        "敏感凭据不进入 Git、文档、截图和普通日志；签名与 Token 通过安全通道注入。",
        "旧功能矩阵应标注为 2026-07-29 前状态，以本实施书和配套进度表为新基线。",
    ):
        add_bullet(doc, text)

    doc.add_heading("附录 A｜资料来源", level=1)
    source_rows = []
    for source in data["sources"]:
        source_rows.append([source["id"], source["name"], source["type"], source["path"], source["used_for"]])
    add_table(doc, ["ID", "资料", "类型", "路径", "用途"], source_rows, [0.4, 1.3, 0.8, 2.6, 1.4], font_size=7.2)

    doc.add_heading("附录 B｜交付与维护说明", level=1)
    add_paragraph(doc, "配套进度表中的进度、风险评分和汇总均由公式计算；后续更新任务状态、完成百分比和权重后，总览会自动刷新。文档与进度表使用同一份 2026-08-01 数据基线。")
    add_paragraph(doc, "本次仅生成交付文档和进度表，不修改 birdpart2 业务源码。", color=MUTED)

    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_docx(load_data())
    print(OUTPUT_PATH)
