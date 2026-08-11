from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\Androidstudio2\project\1\.codex-tmp\项目1任务模块与筛选搜索改善实施批次说明书_2026-08-10.docx")

INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GREEN = "176B5B"
PALE_GREEN = "EAF4F0"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F2F4F7"
MUTED = "5B6573"
RED = "9B1C1C"
PALE_RED = "FCEBE9"
GOLD = "7A5A00"
PALE_GOLD = "FFF6DA"
WHITE = "FFFFFF"
BLACK = "1F2328"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, *, size: float | None = None, color: str | None = None,
                 bold: bool | None = None, italic: bool | None = None,
                 latin: str = "Calibri", east_asia: str = "Microsoft YaHei"):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_margins = tcpr.first_child_found_in("w:tcMar")
    if tc_margins is None:
        tc_margins = OxmlElement("w:tcMar")
        tcpr.append(tc_margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Sequence[int], indent: int = TABLE_INDENT_DXA):
    assert sum(widths) == TABLE_WIDTH_DXA, (widths, sum(widths))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout"):
        existing = tblpr.find(qn(tag))
        if existing is not None:
            tblpr.remove(existing)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tblw.set(qn("w:type"), "dxa")
    tblpr.append(tblw)
    tblind = OxmlElement("w:tblInd")
    tblind.set(qn("w:w"), str(indent))
    tblind.set(qn("w:type"), "dxa")
    tblpr.append(tblind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[idx]
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(width))
            tcw.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def prevent_row_split(row):
    trpr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    trpr.append(cant_split)


def set_paragraph_shading(paragraph, fill: str):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, *, color: str = BLUE, size: int = 8, space: int = 7):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)


def add_page_field(paragraph):
    paragraph.add_run("第 ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])
    paragraph.add_run(" 页")
    for r in paragraph.runs:
        set_run_font(r, size=9, color=MUTED)


def add_custom_numbering(doc: Document):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs or [0]) + 1
    next_num = max(existing_num or [0]) + 1

    def make(kind: str, marker: str, abs_id: int, num_id: int):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        nsid = OxmlElement("w:nsid")
        nsid.set(qn("w:val"), f"{abs_id:08X}"[-8:])
        abstract.append(nsid)
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        numfmt = OxmlElement("w:numFmt")
        numfmt.set(qn("w:val"), kind)
        lvl.append(numfmt)
        lvltext = OxmlElement("w:lvlText")
        lvltext.set(qn("w:val"), marker)
        lvl.append(lvltext)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        ppr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        ppr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.append(spacing)
        lvl.append(ppr)
        if kind == "bullet":
            rpr = OxmlElement("w:rPr")
            rfonts = OxmlElement("w:rFonts")
            rfonts.set(qn("w:ascii"), "Arial")
            rfonts.set(qn("w:hAnsi"), "Arial")
            rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            rpr.append(rfonts)
            lvl.append(rpr)
        abstract.append(lvl)
        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            numbering.insert(list(numbering).index(first_num), abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        aid = OxmlElement("w:abstractNumId")
        aid.set(qn("w:val"), str(abs_id))
        num.append(aid)
        numbering.append(num)

    make("bullet", "•", next_abs, next_num)
    make("decimal", "%1.", next_abs + 1, next_num + 1)
    return next_abs, next_abs + 1


def new_num_instance(doc: Document, abstract_id: int) -> int:
    numbering = doc.part.numbering_part.element
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    num_id = max(existing_num or [0]) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    aid = OxmlElement("w:abstractNumId")
    aid.set(qn("w:val"), str(abstract_id))
    num.append(aid)
    lvl_override = OxmlElement("w:lvlOverride")
    lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    lvl_override.append(start_override)
    num.append(lvl_override)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, num])
    ppr.append(numpr)


def setup_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        s = styles[name]
        s.font.name = "Calibri"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = rgb(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

    if "Table Text" not in styles:
        table_style = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_style = styles["Table Text"]
    table_style.font.name = "Calibri"
    table_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    table_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    table_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    table_style.font.size = Pt(9.5)
    table_style.font.color.rgb = rgb(BLACK)
    table_style.paragraph_format.space_before = Pt(0)
    table_style.paragraph_format.space_after = Pt(2)
    table_style.paragraph_format.line_spacing = 1.08

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Callout"]
    callout.font.name = "Calibri"
    callout._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    callout._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    callout._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    callout.font.size = Pt(10.5)
    callout.font.color.rgb = rgb(INK)
    callout.paragraph_format.left_indent = Inches(0.12)
    callout.paragraph_format.right_indent = Inches(0.12)
    callout.paragraph_format.space_before = Pt(5)
    callout.paragraph_format.space_after = Pt(8)
    callout.paragraph_format.line_spacing = 1.12


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def add_para(doc, text: str, *, bold_lead: str | None = None, color: str | None = None,
             after: float | None = None, italic: bool = False):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True, color=color or BLACK)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, italic=italic, color=color or BLACK)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic, color=color or BLACK)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def add_bullets(doc, items: Iterable[str], bullet_abstract_id: int):
    bullet_id = new_num_instance(doc, bullet_abstract_id)
    for text in items:
        p = doc.add_paragraph()
        apply_num(p, bullet_id)
        r = p.add_run(text)
        set_run_font(r)


def add_steps(doc, items: Iterable[str], decimal_abstract_id: int):
    decimal_id = new_num_instance(doc, decimal_abstract_id)
    for text in items:
        p = doc.add_paragraph()
        apply_num(p, decimal_id)
        r = p.add_run(text)
        set_run_font(r)


def add_callout(doc, label: str, text: str, *, fill: str = PALE_BLUE, accent: str = BLUE):
    p = doc.add_paragraph(style="Callout")
    set_paragraph_shading(p, fill)
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    pbdr.append(left)
    ppr.append(pbdr)
    r1 = p.add_run(f"{label}  ")
    set_run_font(r1, bold=True, color=accent)
    r2 = p.add_run(text)
    set_run_font(r2, color=INK)
    return p


def add_table(doc, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[int],
              *, header_fill: str = PALE_GRAY, compact: bool = False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    prevent_row_split(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.style = doc.styles["Table Text"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=9.5 if compact else 10, bold=True, color=INK)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            p = cell.paragraphs[0]
            p.style = doc.styles["Table Text"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            set_run_font(r, size=9.2 if compact else 9.5)
        set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_batch(doc, code: str, title: str, objective: str, files: str, actions: Sequence[str],
              acceptance: Sequence[str], estimate: str, bullet_id: int):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    badge = p.add_run(f"{code}  ")
    set_run_font(badge, size=12.5, bold=True, color=WHITE)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), GREEN)
    badge._element.get_or_add_rPr().append(shd)
    title_run = p.add_run(title)
    set_run_font(title_run, size=13, bold=True, color=INK)
    add_para(doc, f"目标：{objective}", bold_lead="目标：")
    add_para(doc, f"主要文件：{files}", bold_lead="主要文件：", color=MUTED)
    add_bullets(doc, actions, bullet_id)
    add_para(doc, "完成门槛：", bold_lead="完成门槛：", after=2)
    add_bullets(doc, acceptance, bullet_id)
    add_para(doc, f"估算：{estimate}", bold_lead="估算：", color=GREEN, after=8)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)
    bullet_id, decimal_id = add_custom_numbering(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("项目 1 · 任务模块与筛选搜索改善")
    set_run_font(hr, size=9, color=MUTED, bold=True)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(fp)

    # First-page memo masthead.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("IMPLEMENTATION BRIEF / 2026-08-10")
    set_run_font(r, size=9.5, color=GREEN, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(2)
    title.paragraph_format.space_after = Pt(6)
    title.paragraph_format.keep_with_next = True
    tr = title.add_run("项目 1 任务模块与筛选搜索改善")
    set_run_font(tr, size=23, color=INK, bold=True)
    tr.add_break()
    tr2 = title.add_run("实施批次说明书")
    set_run_font(tr2, size=23, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    sr = subtitle.add_run("基于任务设计稿、现有生产代码与 birdbox-v1 冻结契约的实施基线")
    set_run_font(sr, size=12.5, color=MUTED)

    meta = [
        ("文档编号", "TASK-FILTER-IMPL-20260810"),
        ("版本 / 日期", "1.0 / 2026-08-10"),
        ("适用范围", "任务 Tab、SD 卡/批次、AI、复制、同步、任务详情、报告、相册搜索与筛选"),
        ("契约状态", "birdbox-v1@1.0.0 强制冻结；本批次必需新增外部接口 0 个"),
        ("交付状态", "实施批次基线；需通过真机、视觉与完整自动化门禁后发布"),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f"{label}：")
        set_run_font(r1, size=10.5, bold=True, color=BLACK)
        r2 = p.add_run(value)
        set_run_font(r2, size=10.5, color=BLACK)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    set_paragraph_border(rule, color=GREEN, size=12)

    add_callout(
        doc,
        "执行结论",
        "任务模块不应重做：当前工作树已具备生产入口、SD 卡状态、AI、复制、同步、任务控制与报告主链。"
        "本次优先修复相册搜索/筛选的正确性，再做任务设计收口和发布验收。现有外部接口、参数、字段、状态码与语义一律不变。",
        fill=PALE_GREEN,
        accent=GREEN,
    )
    add_table(
        doc,
        ["判断项", "结论", "实施含义"],
        [
            ["任务模块", "主体已接入，进入收口阶段", "复核设计 01-16、异常恢复、真实盒子与视觉门禁"],
            ["筛选/搜索", "存在明确正确性缺陷", "优先完成“白鹭”等跨页鸟种搜索与组合筛选"],
            ["外部接口", "必需新增 0 个", "只使用冻结 OpenAPI 已声明的请求；未声明参数不得发送"],
            ["发布判断", "当前不放行", "定向测试在本次审计窗口内超时，且缺真机与全量视觉证据"],
        ],
        [1800, 2300, 5260],
        header_fill=PALE_GREEN,
    )

    add_heading(doc, "1. 审计依据与边界", 1)
    add_para(doc, "本说明书基于以下材料完成只读审计，并以当前工作树而不是 8 月 6 日旧结论作为现状依据。")
    add_bullets(doc, [
        r"设计基准：C:\Users\32714\Desktop\v1\overview\task-module-overview.png；task-module 目录下 01-16 状态设计图。",
        r"历史参考：C:\Users\32714\Desktop\v1\任务部分未实现 中的页面清单、旧实施说明书和接口评估说明书。",
        r"当前实现：D:\Androidstudio2\project\1\lib\bird_companion\features\tasks、gallery、copy、jobs、storage 与相关测试。",
        r"冻结权威：docs/contracts/birdbox-v1.openapi.yaml、birdbox-v1-freeze.json、Schema、ADR 与契约测试。",
        "验证限制：本次尝试运行任务与相册定向测试，120 秒内无输出并超时终止；不视为通过或失败，必须在发布批次重新执行。",
    ], bullet_id)
    add_callout(
        doc,
        "冻结优先级",
        "当早期文字版接口说明与冻结 OpenAPI 不一致时，以带 SHA-256 清单的 birdbox-v1 冻结基线为准。"
        "例如照片列表的冻结查询参数并不包含 search、species、tags 等字段，因此不能把这些字段直接加入网络请求。",
        fill=PALE_GOLD,
        accent=GOLD,
    )
    add_callout(
        doc,
        "工作树修改授权",
        "已跟踪但尚未提交的文件，以及 Git 未追踪文件，只要与任务板块或筛选/搜索问题直接相关，均允许在本批次修改、纳入测试并提交；"
        "无关文件不得顺带整理、删除或覆盖，也不以清理工作树为实施前提。",
        fill=PALE_BLUE,
        accent=BLUE,
    )

    add_heading(doc, "2. 项目 1 当前任务模块审计", 1)
    add_para(doc, "8 月 6 日说明中的 B1-B6 已有大量实现落在当前工作树。后续工作应以“验证、收口、修正意图偏差”为主，避免再次建立平行页面或 Demo 链路。")
    add_table(
        doc,
        ["设计范围", "当前生产实现", "判断", "本次处理"],
        [
            ["overview / 01 任务首页与可执行操作", "TaskExperienceRoot、TaskHomePage、TaskHomeCapabilityResolver", "已接生产数据", "复核可用条件、禁用原因、防重复点击与返回栈"],
            ["02-05 SD 卡状态", "SdCardFlowPage + scan/rescan", "已覆盖主要状态", "补真机卡拔插、窄屏/字体放大与错误文案验收"],
            ["06-07、10、12-14 任务详情", "TaskDetailPage + available_actions/version + 事件/轮询", "主链已实现", "逐状态核对按钮、冲突、断线与取消后不自动返回"],
            ["08-09 复制步骤", "生产 CopyConfirmationFlow + Cubit + estimate/create", "已接真实接口", "复核目标盘、空间、二次估算、409 与幂等"],
            ["11 同步结果", "BirdSyncService + TaskSyncResultSheet", "功能已实现，形态有意偏离", "保持 App 本地同步结果，不伪造盒子任务"],
            ["15-16 复制结果/报告", "ProductionTaskResultPage + JobReport + failures 分页", "主体已实现", "缺失权威字段时隐藏或标“未提供”，禁止造数"],
            ["跨页/重连", "RepositoryTaskExperienceController、EventClient、Session", "已有恢复逻辑", "真机验证最终一致、旧会话隔离与重复导航"],
        ],
        [1600, 2600, 1700, 3460],
        compact=True,
    )
    add_callout(
        doc,
        "实施原则",
        "任务状态只由 BirdJobStatus、available_actions、version 和 JobReport 驱动；App 不根据百分比推断状态，不为满足截图自行补造 XMP 数量或同步任务。",
        fill=PALE_BLUE,
        accent=BLUE,
    )

    add_heading(doc, "3. 筛选与搜索失效的根因", 1)
    add_para(doc, "用户输入“白鹭”无结果不是输入框或按钮失灵，而是查询对象、网络请求和本地过滤三层没有形成闭环。")
    add_steps(doc, [
        "FilterSheet 和 GallerySearchDialog 会把搜索词、鸟种、标签、评分、识别度、清晰度、识别结果和推荐条件写入 PhotoQuery。",
        "GalleryCubit.search/refresh 会把 PhotoQuery 交给 PhotoRepository.page，并具备 350 ms 防抖和 request generation 旧请求保护。",
        "PhotoApi.page 实际只发送 PhotoQuery.parameters；该集合当前只包含 page_size、sort、cursor、keep_state、analysis_state、group_id、scene_id。",
        "在线成功路径直接返回服务器当前页，没有执行 PhotoRepositoryImpl._matches；_matches 仅在 API 异常并回退到已有缓存快照时使用。",
        "因此 search/species/tags/min_score/min_confidence/clarity_state/recognition_state/recommended_only 在在线查询中既没有发送，也没有本地执行；匹配照片若位于后续页也不会被发现。",
    ], decimal_id)
    add_table(
        doc,
        ["层", "已有能力", "断点", "影响"],
        [
            ["UI", "字段齐全、搜索按钮与筛选面板存在", "不知道哪些条件由服务端执行", "表面可输入，实际无结果"],
            ["Query", "toJson 可保存完整筛选状态", "parameters 只含冻结子集", "恢复状态与真实请求语义不一致"],
            ["Repository", "已有 _matches 和本地缓存回退", "在线成功时不调用；快照可能不完整", "“白鹭”只在偶然缓存命中时可见"],
            ["分页", "支持 cursor/page_size", "本地筛选缺少独立游标和完整性标记", "后续页匹配项遗漏或重复"],
        ],
        [1400, 2600, 2500, 2860],
        header_fill=PALE_RED,
    )
    add_callout(
        doc,
        "不允许的快速修复",
        "不得直接把 search、species、tags 等加入 GET /projects/{id}/files 的查询串，因为冻结 OpenAPI 未声明这些参数。"
        "这样做会造成客户端与真实盒子契约漂移，即使 mock server 当前能够识别也不能作为发布依据。",
        fill=PALE_RED,
        accent=RED,
    )

    add_heading(doc, "4. 目标实现：冻结接口下的完整筛选", 1)
    add_para(doc, "目标链路如下：")
    flow = doc.add_paragraph()
    flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    flow.paragraph_format.space_before = Pt(6)
    flow.paragraph_format.space_after = Pt(10)
    fr = flow.add_run("搜索/筛选 UI  ->  QueryPlan  ->  冻结服务端筛选+分页  ->  本地完整谓词  ->  本地结果会话/游标  ->  相册")
    set_run_font(fr, size=10.5, bold=True, color=GREEN)
    add_para(doc, "QueryPlan 将条件显式分成两组：")
    add_table(
        doc,
        ["执行位置", "条件", "处理方式"],
        [
            ["盒子端（冻结允许）", "page_size、cursor、sort、keep_state、analysis_state、group_id、scene_id", "继续由 PhotoApi 发送；page_size 可在完整扫描时提高到冻结上限 200"],
            ["App 本地", "search、species、min_score、min_confidence、tags、clarity_state、recognition_state、recommended_only", "拉取候选集后复用/抽取 _matches；不出现在网络请求中"],
            ["鸟种辅助", "中文名、species_id、英文名、拉丁名", "复用 GET /api/v1/species 做别名解析，再与 PhotoSummary.recognition.candidates 比对"],
        ],
        [1900, 3500, 3960],
        header_fill=PALE_GREEN,
    )
    add_heading(doc, "4.1 本地结果会话", 2)
    add_bullets(doc, [
        "当存在 App 本地条件时，按冻结服务端条件逐页拉取候选集，直到 has_more=false；每页到达后更新“已扫描 N 张”进度。",
        "完整候选集按 device namespace + project_id + 服务端条件 + sort 生成会话键；保存去重后的 PhotoSummary 和 complete 标记。",
        "本地游标使用内部前缀（例如 local-v1:<session>:<offset>），只在 Repository/Cubit 内消费，绝不能传入 PhotoApi。",
        "搜索词变化时增加 generation，旧会话结果即使晚到也不能覆盖新查询；复用当前 GalleryCubit 的旧请求保护并补充扫描任务取消/去重。",
        "照片、审阅、任务完成或批次变化事件到达时失效相关筛选会话；仅设置筛选条件变化时复用未过期候选集。",
        "离线且快照不完整时显示“仅在已缓存照片中查找”，不得把“缓存中无结果”表述为“项目中无结果”。",
    ], bullet_id)
    add_heading(doc, "4.2 匹配语义", 2)
    add_table(
        doc,
        ["条件", "匹配规则", "示例验收"],
        [
            ["通用搜索", "文件名、用户标签、候选鸟种中文/英文/拉丁名包含匹配；去首尾空格，英文忽略大小写", "“白鹭”、little egret、DSC_1820、湿地均可命中"],
            ["鸟种", "species_id 或四类名称精确/包含匹配；候选列表任一项可命中", "little-egret 与“白鹭”指向同一照片集合"],
            ["标签", "默认 AND：所选标签必须全部存在；比较前统一 trim 与大小写", "湿地+晨拍只返回同时具备两标签的照片"],
            ["评分/识别度", "评分取 total_score；识别度取首候选 confidence；缺值不满足最小阈值", "≥4 星且 ≥85% 不会把缺分析照片纳入"],
            ["识别结果", "recognized / needs_review / unknown 与 candidates、low_confidence 一致", "低置信度照片只进入 needs_review"],
            ["组合", "不同维度使用 AND；同一鸟种候选为 OR", "白鹭 + 保留 + 清晰 + ≥85% 同时满足"],
        ],
        [1500, 4860, 3000],
        compact=True,
    )

    add_heading(doc, "5. 本次实施批次（B8-B12）", 1)
    add_callout(doc, "推荐顺序", "主链 B8 -> B9 -> B10 -> B12；B11 在 B8 后可与 B9/B10 并行，最终在 B12 汇合。总估算 8-13 人日，不含盒子端开发，因为本批次不改盒子接口。", fill=PALE_GREEN, accent=GREEN)

    add_batch(
        doc,
        "B8",
        "冻结刷新、现状快照与失败复现",
        "登记当前工作树、冻结接口和可复现缺陷；允许修改与本问题相关的未提交及未追踪文件，同时避免覆盖无关改动。",
        "docs/contracts/*、photo_query.dart、photo_api.dart、现有任务/相册测试",
        [
            "保存当前变更清单和基线 SHA；已跟踪但未提交、以及 Git 未追踪的文件，只要与任务或筛选问题直接相关均可修改；无关文件保持不动。",
            "增加失败用例：在线模式搜索“白鹭”，匹配照片位于第 2 页或更后；断言当前实现返回空。",
            "增加网络契约断言：照片列表请求不得携带冻结 OpenAPI 未声明参数。",
            "把 01-16 设计图逐项映射到当前生产文件，并标记“已实现/需收口/有意偏离”。",
        ],
        [
            "缺陷可稳定复现，测试能防止把未声明参数偷加到请求中。",
            "现有 B1-B6 代码保持可追溯；问题相关的未提交/未追踪文件纳入本批次，无关用户改动保持不动。",
        ],
        "0.5-1 人日",
        bullet_id,
    )

    add_batch(
        doc,
        "B9",
        "搜索/筛选正确性闭环",
        "在零外部接口变更下，让“白鹭”和全部已有筛选项真正影响在线结果。",
        "photo_query.dart、photo_repository.dart、photo_repository_impl.dart、photo_api.dart、gallery_cubit.dart",
        [
            "新增内部 QueryPlan：serverParameters 与 localPredicate 明确分离；保留现有 q.parameters 对冻结请求的约束。",
            "把 _matches 抽成可单测的 PhotoLocalFilterEngine；统一中文、英文、拉丁名、species_id 与标签匹配。",
            "当存在本地条件时逐页获取候选照片，使用 page_size=200，去重、过滤并构建完整结果会话。",
            "为本地结果建立独立游标；严禁 local cursor 或本地条件进入 PhotoApi。",
            "复用 GET /species 解析鸟种名称和别名，查询失败时仍可直接按照片候选名称匹配。",
        ],
        [
            "在线搜索白鹭可命中第 2 页及以后照片；清空搜索恢复完整列表。",
            "所有单条件与组合条件都在真实在线成功路径生效。",
            "抓包/Mock 断言显示外发参数集合与冻结 OpenAPI 完全一致。",
        ],
        "2-3 人日",
        bullet_id,
    )

    add_batch(
        doc,
        "B10",
        "分页、缓存、性能与筛选体验",
        "让完整筛选在 3,000-10,000 张照片规模下可解释、可取消、可恢复。",
        "gallery_cubit.dart、gallery_page.dart、filter_sheet.dart、gallery_search_dialog.dart、local_cache.dart",
        [
            "GalleryState 增加 filtering/scannedCount/resultComplete/cacheScope，不用普通 loading 掩盖旧结果。",
            "扫描时显示进度与取消入口；搜索变化立即淘汰旧结果，避免“白鹭”结果被更早请求覆盖。",
            "缓存完整候选集和 complete 标记；照片数据变更、设备切换、批次切换时精准失效。",
            "补充无结果、离线缓存不完整、服务端失败、别名解析失败等文案；重置必须同时清空搜索与全部高级条件。",
            "保持筛选状态返回恢复，但只保存逻辑条件，不持久化短期本地游标。",
        ],
        [
            "3,672 张基准数据下 UI 操作反馈 <100 ms；首批扫描进度可见 <300 ms；完整结果目标 <10 s（真机局域网验收）。",
            "连续输入/快速切换条件不会显示过期结果；离线不会给出误导性“项目无结果”。",
            "加载更多无重复、无漏项，返回详情后仍保持条件和滚动语义。",
        ],
        "2-3 人日",
        bullet_id,
    )

    add_batch(
        doc,
        "B11",
        "任务设计收口与有意偏离确认",
        "以 01-16 设计稿做视觉/交互收口，不建立第二套生产链或伪造盒子数据。",
        "task_experience_root.dart、repository_task_experience_controller.dart、task_home_page.dart、sd_card_flow_page.dart、copy_confirmation_page.dart、task_detail_page.dart、production_task_result_page.dart",
        [
            "首页和 + 弹层共用能力解析；无项目、断线、扫描中、项目任务冲突时展示可解释状态。",
            "SD detected/missing/unreadable/empty 与重扫、设备详情、批次创建入口逐图验收。",
            "复制步骤 2/3 保持提交前重新估算、目标在线/空间校验、version 和幂等；成功后 replacement 到真实任务详情。",
            "详情按钮只来自 available_actions，断线时保留但禁用；取消二次确认后等待权威快照。",
            "复制报告只显示冻结 JobReport 字段；XMP 数量缺失时隐藏或写“盒子未提供”，不得复用照片数冒充。",
            "同步继续使用 BirdSyncService + 本地结果弹层；不把本地同步伪装为持久化 BirdJobStatus。",
        ],
        [
            "01-16 每张设计图均有生产可达页面、明确状态覆盖或书面“有意偏离”决议。",
            "窄屏 360、常规 426、文字 1.3x/1.5x 无裁切、重叠和不可达按钮。",
            "AI、复制、同步、导入四类入口不存在 SnackBar 占位或 Demo ID。",
        ],
        "2-3 人日，可与 B9/B10 并行",
        bullet_id,
    )

    add_batch(
        doc,
        "B12",
        "契约、真机、视觉与发布门禁",
        "用完整证据决定发布，不以“已有代码/已有测试文件”代替通过。",
        "test/bird_companion、test/contracts、tool/contracts、真机记录与发布清单",
        [
            "运行契约 SHA 校验、照片查询参数测试、任务/复制/同步/恢复测试和全量 flutter test。",
            "使用真实盒子验证 3,672+ 照片白鹭跨页搜索、组合筛选、断线/重连、缓存不完整和设备切换。",
            "逐图比对 01-16；覆盖卡拔插、目标盘拔出、低电量、过温、409、401/403 与签名 URL 过期。",
            "记录 App SHA、固件 SHA、契约版本、测试命令/结果、截图、已知偏离与回滚点。",
        ],
        [
            "所有 P0/P1 缺陷关闭；外发接口零漂移；白鹭与其他鸟种搜索通过。",
            "全量自动化、真实盒子和视觉验收均有可追溯证据；否则不发布。",
        ],
        "1.5-3 人日",
        bullet_id,
    )

    add_heading(doc, "6. 接口冻结清单与新增接口评估", 1)
    add_callout(doc, "本批次接口结论", "必需新增外部接口：0 个；修改现有外部接口：0 个。所有实现均为 App 内部查询规划、缓存、状态和页面收口。", fill=PALE_GREEN, accent=GREEN)
    add_table(
        doc,
        ["功能", "冻结接口", "允许的本次使用", "禁止变更"],
        [
            ["照片列表", "GET /api/v1/projects/{projectId}/files", "cursor、page_size、sort、keep_state、analysis_state、group_id、scene_id", "不得新增 search/species/tags/评分等查询参数"],
            ["鸟种辅助", "GET /api/v1/species", "search、page_size；仅用于名称/ID 候选", "不得改变返回字段或把它当照片搜索接口"],
            ["SD/批次", "scan、rescan、projects、imports、current", "继续驱动卡状态与导入入口", "路径、请求体、状态码不变"],
            ["AI", "POST /projects/{id}/analysis-jobs", "手动/自动创建共享幂等保护", "不传模型参数，不新增动作"],
            ["复制", "GET copy/estimate；POST copy", "提交前重估，使用 target_id、version、XMP/校验布尔值", "不硬编码目标，不改字段语义"],
            ["任务", "GET jobs/detail/report/failures；POST actions；WS events", "权威状态、动作、报告、失败分页与恢复", "不推断状态，不扩展 action 枚举"],
            ["同步", "现有 decision 与 files/actions", "重放审阅修改并显示本地同步结果", "不新增伪 sync job，不离线排队任务控制"],
        ],
        [1400, 2550, 3300, 2110],
        compact=True,
    )
    add_heading(doc, "6.1 未来可选提案（本批次不实施）", 2)
    add_para(doc, "以下能力只有在产品坚持相应体验、且 App 侧方案无法满足性能/审计目标时才需要新版本提案。它们不是本批次新增接口。")
    add_table(
        doc,
        ["潜在提案", "触发条件", "可能的契约变化", "当前处理"],
        [
            ["服务端全量筛选", "10,000+ 照片下 App 完整扫描无法满足目标", "新版本照片搜索端点或显式声明筛选参数", "先做 App 本地完整筛选；不改 v1"],
            ["报告 XMP/校验计数", "产品要求设计 15/16 精确显示权威数量", "JobReport 新版本增加 xmp_*、verification_* 字段", "隐藏或标“未提供”；不造数"],
            ["持久化同步任务", "产品要求同步出现在跨设备任务历史与详情", "新增 versioned sync job 创建/查询能力", "使用本地 TaskSyncResultSheet"],
        ],
        [1900, 2700, 2800, 1960],
        compact=True,
    )
    add_callout(doc, "提案规则", "任何未来新增能力必须另立提案、OpenAPI、Schema、fixture、ADR、契约测试、兼容与回滚计划；未批准前不得修改 birdbox-v1-freeze.json 或向 mock/真实盒子发送新字段。", fill=PALE_GOLD, accent=GOLD)

    add_heading(doc, "7. 文件级实施清单", 1)
    add_table(
        doc,
        ["文件/区域", "计划变更", "外部接口影响"],
        [
            ["gallery/domain/photo_query.dart", "增加服务端/本地条件分类、规范化签名和仅内部游标语义", "无"],
            ["gallery/domain/photo_repository.dart", "扩展内部结果元数据或增加过滤会话抽象", "无（Dart 内部接口）"],
            ["gallery/data/photo_repository_impl.dart", "在线成功路径执行完整过滤、跨页候选集、完整性缓存、失效与去重", "无"],
            ["gallery/data/photo_api.dart", "只接受 serverParameters；增加防止本地游标外发的断言", "无；强化冻结"],
            ["gallery/presentation/gallery_cubit.dart", "扫描进度、取消/代际、结果完整性和状态恢复", "无"],
            ["filter_sheet.dart / gallery_search_dialog.dart", "字段语义、无结果/扫描反馈、重置和可访问性", "无"],
            ["tasks/* / copy/* / jobs/*", "按 B11 收口，不新建平行页面；权威字段缺失时显式降级", "无"],
            ["test/* / tool/contracts/*", "跨页白鹭、组合筛选、本地游标、零接口漂移、任务设计/恢复门禁", "无"],
        ],
        [2500, 4930, 1930],
        compact=True,
    )

    add_heading(doc, "8. 测试与验收矩阵", 1)
    add_table(
        doc,
        ["层级", "必须覆盖", "关键断言"],
        [
            ["纯单元", "PhotoLocalFilterEngine、QueryPlan、规范化、排序", "白鹭/英文/拉丁名/ID、标签 AND、缺值阈值、组合条件"],
            ["Repository", "跨页扫描、缓存完整性、本地游标、失效", "第 2 页以后可命中；无重复/漏项；local cursor 不外发"],
            ["API 契约", "照片列表与 species 请求参数", "仅发送 OpenAPI 声明集合；冻结 SHA 不变"],
            ["Cubit", "快速输入、旧请求晚到、加载更多、返回恢复", "最后一次查询胜出；清空恢复；状态不串批次/设备"],
            ["Widget", "搜索、筛选、进度、无结果、离线提示", "按钮可达、反馈明确、360/426 宽与 1.5x 字体可用"],
            ["任务模块", "01-16、available_actions、version、复制/同步/报告", "真实 ID、无 Demo、无造数、断线写操作禁用"],
            ["真机", "3,672+ 照片、弱网、重连、卡/盘拔插、冲突", "结果完整、最终一致、无重复任务和误导空态"],
            ["发布", "全量 flutter test、契约、视觉、SHA 与回滚", "全部证据归档后才放行"],
        ],
        [1500, 3900, 3960],
        compact=True,
    )
    add_heading(doc, "8.1 白鹭专项验收脚本", 2)
    add_steps(doc, [
        "准备至少 3 个分页的数据：第 1 页无白鹭，第 2/3 页各含白鹭；同时包含翠鸟、东方大苇莺和同名标签干扰项。",
        "在搜索入口输入“白鹭”，确认 UI 出现扫描进度，最终只返回鸟种候选或用户标签/文件名包含白鹭的照片。",
        "用鸟种筛选分别输入“白鹭”“little-egret”“Little Egret”，确认结果集合一致。",
        "叠加“保留 + 清晰 + 识别度≥85%”，确认使用 AND；取消任一条件后集合按预期扩大。",
        "抓取请求并确认 /files 只携带冻结参数，/species 只携带 search/page_size；本地游标从未出现在网络层。",
        "断网后重复搜索：若缓存完整，结果一致并标记缓存来源；若缓存不完整，显示范围提示而不是错误宣称无结果。",
    ], decimal_id)

    add_heading(doc, "9. 风险、回滚与发布门禁", 1)
    add_table(
        doc,
        ["风险", "预防措施", "回滚方式"],
        [
            ["跨页扫描请求过多", "page_size=200、候选集缓存、服务端冻结条件预过滤、同查询去重", "关闭本地完整筛选开关并恢复旧分页，不改接口"],
            ["缓存不完整导致假空态", "complete 标记、范围提示、事件精准失效", "禁用离线筛选，仅保留已缓存浏览"],
            ["旧请求覆盖新搜索", "generation + 扫描任务取消/忽略", "退回到单次提交搜索，禁用实时防抖"],
            ["任务模块回归", "只做局部收口，复用现有 Controller/Repository；逐图测试", "按批次提交回退 B11，不影响 B9/B10"],
            ["接口漂移", "冻结 SHA、网络参数白名单、PR 门禁", "拒绝合并；恢复冻结文件并重新跑契约测试"],
            ["相关与无关工作树改动混淆", "B8 分类登记；问题相关的未提交/未追踪文件可修改，无关文件不动；不做 reset/checkout/clean 清理", "只回退本批次相关修改，保留其他用户改动"],
        ],
        [2200, 4440, 2720],
        compact=True,
    )
    add_heading(doc, "9.1 发布前必须全部满足", 2)
    add_bullets(doc, [
        "birdbox-v1 OpenAPI、Schema、ADR、原接口说明和 freeze SHA 无变化；新增外部接口登记仍为 0。",
        "白鹭与其他鸟种跨页搜索、全部筛选项、组合筛选、排序、加载更多和返回恢复通过。",
        "任务设计 01-16 的生产可达性、按键、返回栈、断线/重连、失败与取消状态通过。",
        "XMP/同步等有意偏离获得产品确认；文案不暗示盒子返回了不存在的字段或任务。",
        "完整自动化不再超时且结果归档；真机 App SHA、固件 SHA、契约版本和回滚点齐全。",
    ], bullet_id)

    add_heading(doc, "10. 最终实施决议", 1)
    add_callout(
        doc,
        "批准建议",
        "批准 B8-B12 作为 8 月 10 日改善批次：先解决筛选/搜索正确性，再收口任务设计并完成发布门禁。"
        "本批次不修改任何现有外部接口，也不新增外部接口；如未来要求服务端全筛选、权威 XMP 计数或持久化同步任务，必须另行提案并及时通知接口 Owner。",
        fill=PALE_GREEN,
        accent=GREEN,
    )
    # Metadata and core properties.
    props = doc.core_properties
    props.title = "项目1任务模块与筛选搜索改善实施批次说明书"
    props.subject = "任务模块设计收口、筛选搜索修复、接口冻结与发布门禁"
    props.author = "Codex"
    props.keywords = "项目1, 任务模块, 白鹭, 搜索, 筛选, birdbox-v1, 接口冻结"
    props.comments = "Generated from the 2026-08-10 read-only implementation audit."

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
